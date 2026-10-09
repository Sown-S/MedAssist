"""API xem / xuất nhật ký cho Admin (US09)."""
import csv
import io
import uuid

import pytest
from openpyxl import load_workbook


@pytest.fixture
def admin(login):
    return login("admin")


def test_only_admin_can_read_logs(client, users, login, logs):
    for who in ("bs", "yta"):
        assert client.get("/api/v1/admin/audit-logs", headers=login(who)).status_code == 403
    denied = logs(action="VIEW", entity="audit_log")
    assert len(denied) == 2 and all(r.new_values["result"] == "denied" for r in denied)
    assert client.get("/api/v1/admin/audit-logs").status_code == 401


def test_list_filters_and_paging(client, users, admin):
    tag = f"ma_{uuid.uuid4().hex[:8]}"
    for _ in range(3):
        client.post("/api/v1/auth/login", json={"username": tag, "password": "x"})

    page = client.get("/api/v1/admin/audit-logs", headers=admin,
                      params={"username": tag, "action_type": "LOGIN_FAILED", "limit": 2}).json()
    assert page["total"] == 3 and len(page["items"]) == 2 and page["limit"] == 2
    item = page["items"][0]
    assert item["attempted_username"] == tag and item["username"] is None
    assert item["new_values"] == {"reason": "unknown_user"}

    rest = client.get("/api/v1/admin/audit-logs", headers=admin,
                      params={"username": tag, "limit": 2, "offset": 2}).json()
    assert len(rest["items"]) == 1
    assert {i["log_id"] for i in page["items"]}.isdisjoint({i["log_id"] for i in rest["items"]})


def test_list_newest_first_and_joins_username(client, users, admin):
    items = client.get("/api/v1/admin/audit-logs", headers=admin, params={"limit": 5}).json()["items"]
    times = [i["created_at"] for i in items]
    assert times == sorted(times, reverse=True)
    assert items[0]["action_type"] == "VIEW" and items[0]["username"] == "admin"   # chính lượt xem này


@pytest.mark.parametrize("params", [
    {"action_type": "HACK"},
    {"user_id": "khong-phai-uuid"},
    {"limit": 0}, {"limit": 201}, {"offset": -1},
    {"created_from": "2026-10-11T00:00:00+07:00", "created_to": "2026-10-10T00:00:00+07:00"},
])
def test_list_rejects_bad_params(client, users, admin, params):
    assert client.get("/api/v1/admin/audit-logs", headers=admin, params=params).status_code == 422


def test_export_xlsx_is_readable_and_formula_safe(client, users, admin, logs):
    evil = '=HYPERLINK("http://evil.example","Bấm")'
    client.post("/api/v1/auth/login", json={"username": evil, "password": "x"})
    client.post("/api/v1/auth/login", json={"username": "Y tá thử", "password": "x"})

    r = client.get("/api/v1/admin/audit-logs/export", headers=admin,
                   params={"action_type": "LOGIN_FAILED"})
    assert r.status_code == 200
    assert r.headers["content-type"].startswith("application/vnd.openxmlformats")
    assert r.headers["content-disposition"].endswith('.xlsx"')

    ws = load_workbook(io.BytesIO(r.content)).active
    header = [c.value for c in ws[1]]
    assert header[:4] == ["created_at", "action_type", "username", "attempted_username"]
    cells = {row[3].value: row[3] for row in ws.iter_rows(min_row=2)}
    assert "Y tá thử" in cells
    assert cells[evil].data_type == "s"                         # chữ, không phải công thức

    (exp,) = logs(action="EXPORT")
    assert exp.new_values["format"] == "xlsx" and exp.new_values["filters"] == {"action_type": "LOGIN_FAILED"}
    assert exp.new_values["rows"] == ws.max_row - 1 and exp.new_values["truncated"] is False


def test_export_csv_has_bom_and_escapes_formulas(client, users, admin):
    client.post("/api/v1/auth/login", json={"username": "=1+1", "password": "x"})
    r = client.get("/api/v1/admin/audit-logs/export", headers=admin, params={"format": "csv"})
    assert r.status_code == 200 and r.content[:3] == b"\xef\xbb\xbf"
    rows = list(csv.reader(io.StringIO(r.content.decode("utf-8-sig"))))
    attempted = [row[3] for row in rows[1:]]
    assert "'=1+1" in attempted and "=1+1" not in attempted


def test_export_requires_admin(client, users, login):
    assert client.get("/api/v1/admin/audit-logs/export", headers=login("bs")).status_code == 403
