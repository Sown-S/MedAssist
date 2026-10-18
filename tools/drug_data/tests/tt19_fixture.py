"""Bản dựng lại một phần Phụ lục I TT 19/2018/TT-BYT theo đúng cấu trúc bảng của file trên vbpl.moj.gov.vn
(cột TT lặp hai lần, dòng tiêu đề nhóm gộp ngang, dòng tiếp theo bỏ trống STT/tên/đường dùng, lỗi chính tả
giữ nguyên như văn bản gốc). Chỉ gồm các STT mà data/drugs.csv tham chiếu + vài STT để thử bộ lọc."""
from docx import Document

S = "S"
ROWS = [
    (S, "II. THUỐC GIẢM ĐAU VÀ CHĂM SÓC GIẢM NHẸ"),
    (S, "1. Thuốc giảm đau không opi, hạ sốt, chống viêm không steroid"),
    ("15", "Diclofenac", "Uống", "Viên", "25mg, 50mg, 75mg, 100mg"),
    ("16", "Ibuprofen (*)", "Uống", "Viên", "200mg, 400mg"), ("", "", "", "Dung dịch", "40mg/ml"),
    ("17", "Meloxicam", "Uống", "Viên", "7,5mg, 15mg"), ("", "", "Tiêm", "Dung dịch", "10mg/ml"),
    ("18", "Paracetamol (*)", "Uống", "Viên", "80mg, 100mg, 250mg, 500mg"),
    ("", "", "", "Bột pha dung dịch", "80mg, 100mg, 250mg, 500mg"),
    ("", "", "Đặt trực tràng", "Viên đạn", "80mg, 150mg, 300mg"),
    (S, "2. Thuốc giảm đau loại opi"),
    ("21", "Morphin Sulfat", "Uống", "Viên", "10mg, 30mg"),
    (S, "3. Thuốc chăm sóc giảm nhẹ"),
    ("31", "Lactulose (*)", "Uống", "Dung dịch, Sirô", "Từ 0,62 đến 0,74g/ml"),
    (S, "III. THUỐC CHỐNG DỊ ỨNG VÀ DÙNG TRONG CÁC TRƯỜNG HỢP QUÁ MẪN"),
    ("34", "Alimemazin (*)", "Uống", "Viên", "5mg"), ("", "", "", "Sirô", "5mg/ml"),
    ("35", "Clorpheniramin maleat", "Uống (*)", "Viên", "4mg"), ("", "", "", "Sirô", "0,4mg/ml"),
    ("", "", "Tiêm", "Dung dịch", "10mg/ml"),
    ("36", "Dexamethason", "Uống", "Viên", "0,5mg, 1mg"),
    ("", "Dexamethason phosphat (natri)", "Tiêm", "Dung dịch", "4mg/ml"),
    ("40", "Prednisolon", "uống", "Viên", "5mg"), ("", "", "", "Sirô", "5mg/ml"),
    ("41", "Loratadin (*)", "Uống", "Viên", "10mg"),
    ("42", "Promethazin hydroclorid (*)", "Uống", "Viên nén", "10mg, 50mg"),
    (S, "VI. THUỐC TRỊ KÝ SINH TRÙNG, CHỐNG NHIỄM KHUẨN"), (S, "1. Thuốc trị giun, sán"),
    (S, "a. Thuốc trị giun, sán đường ruột"),
    ("66", "Albendazol(*)", "Uống", "Viên", "200mg, 400mg"), ("67", "Mebendazol (*)", "Uống", "Viên", "100mg, 500mg"),
    (S, "2. Thuốc chống nhiễm khuẩn"), (S, "a. Thuốc nhóm beta- lactam"),
    ("77", "Amoxicilin", "Uống", "Viên", "250mg, 500mg"), ("", "", "", "Bột pha hỗn dịch", "125mg, 250mg"),
    ("", "", "Tiêm", "Bột pha tiêm", "500mg, 1g"),
    ("78", "Amoxicilin + Acid Clavulanic", "Uống", "Viên", "500mg + 125mg, 250mg + 62,5mg, 125mg +31,25mg"),
    ("", "", "", "Bột pha hỗn dịch", "500mg + 125mg/5ml, 250mg + 62,5mg/5ml, 125mg + 31,25mg/5ml (sau pha)"),
    ("", "", "Tiêm", "Bột pha tiêm", "500mg + 100mg, 1g + 200mg"),
    ("83", "Cephalexin", "Uống", "Bột pha hỗn dịch", "125mg/5ml, 250mg/5ml (sau pha)"),
    ("", "", "", "Viên", "125mg, 250mg, 500mg"),
    ("85", "Cefixim", "Uống", "Viên", "100mg, 200mg"),
    ("88", "Cefuroxim (axetil)", "Uống", "Viên", "250mg, 500mg"),
    ("", "Cefuroxim (natri)", "Tiêm", "Bột pha tiêm", "750mg, 1,5g"),
    (S, "b. Các thuốc kháng khuẩn khác"), (S, "♦ Thuốc nhóm Nitroimidazol"),
    ("97", "Metronidazol", "Uống", "Viên", "250mg, 500mg"), ("", "", "Trực tràng", "Viên đạn", "500mg, 1g"),
    (S, "♦ Thuốc nhóm macrolid"),
    ("99", "Azithromycin", "Uống", "Viên", "500mg"), ("", "", "", "Bột pha hỗn dịch", "200mg/5ml (sau pha)"),
    ("100", "Clarithromycin", "Uống", "Viên", "250mg, 500mg"),
    (S, "♦ Thuốc nhóm quinolon"),
    ("102", "Ciprofloxacin (hydroclorid)", "Uống", "Viên", "250mg, 500mg"),
    ("103", "Levofloxacin", "Uống", "Viên nén bao phim", "250mg, 500mg, 750mg"),
    (S, "♦ Thuốc nhóm sulfamid"),
    ("104", "Sulfamethoxazol + Trimethoprim", "Uống", "Viên", "800mg + 160mg, 400mg +80 mg, 100mg + 20mg"),
    ("", "", "", "Hỗn dịch", "40mg + 8mg/ml"),
    (S, "♦ Thuốc nhóm nitrofuran"), ("105", "Nitrofurantoin", "Uống", "Viên", "100mg"),
    (S, "♦ Thuốc nhóm tetracyclin"), ("106", "Doxycyclin (hydroclorid)", "Uống", "Viên", "50mg, 100mg"),
    (S, "c. Thuốc chữa bệnh phong"), ("109", "Rifampicin", "Uống", "Viên", "150mg, 300mg"),
    (S, "3. Thuốc chống nấm"), ("130", "Fluconazol", "Uống", "Viên", "50mg, 150mg"),
    (S, "5. Thuốc chống virus"), (S, "a. Thuốc chống virus herpes"),
    ("157", "Aciclovir", "Uống", "Viên", "200mg, 800mg"), ("", "", "", "Hỗn dịch", "40mg/ml"),
    ("", "", "Dùng ngoài (*)", "Kem", "5%"),
    (S, "đ. Thuốc khác"), ("172", "Oseltamivir", "Uống", "Viên", "30mg, 45mg, 75mg"),
    (S, "X. THUỐC TÁC DỤNG ĐỐI VỚI MÁU"), (S, "1. Thuốc chống thiếu máu"),
    ("211", "Acid Folic (*)", "Uống", "Viên", "1mg, 5mg"),
    ("214", "Sắt (sulfat + acid folic) (*)", "Uống", "Viên", "60mg + 0,25mg"),
    (S, "XII. THUỐC TIM MẠCH"), (S, "1. Thuốc chống đau thắt ngực"),
    ("231", "Atenolol", "Uống", "Viên", "50mg, 100mg"),
    (S, "2. Thuốc chống loạn nhịp"), ("238", "Propranolol hydroclorid", "Uống", "Viên", "40mg"),
    (S, "3. Thuốc điều trị tăng huyết áp"),
    ("242", "Amlodipin", "Uống", "Viên", "2,5 mg, 5mg"), ("243", "Captopril", "Uống", "Viên", "25mg, 50mg"),
    ("244", "Enalapril", "Uống", "Viên", "5mg, 10mg, 20mg"), ("245", "Furosemid", "Uống", "Viên", "40mg"),
    ("", "", "Tiêm", "Dung dịch", "10mg/ml"),
    ("246", "Hydroclorothiazid", "Uống", "Viên", "25mg, 50mg"),
    ("248", "Nifedipin", "Uống", "Viên", "5mg, 10mg"), ("", "", "", "Viên tác dụng chậm", "20mg"),
    (S, "5. Thuốc điều trị suy tim"), ("260", "Spironolacton", "Uống", "Viên nén", "25mg"),
    (S, "6. Thuốc chống huyết khối"), ("261", "Acid Acetylsalicylic (*)", "Uống", "Viên", "100mg, 81mg"),
    (S, "7. Thuốc hạ lipit máu"),
    ("263", "Atorvastatin", "Uống", "Viên", "10mg, 20mg"), ("264", "Fenofibrat", "Uống", "Viên", "100mg, 300mg"),
    ("265", "Simvastatin", "Uống", "Viên", "5mg, 10mg, 20mg, 40mg"),
    (S, "XIII. THUỐC NGOÀI DA"), (S, "1. Thuốc chống nấm"),
    ("269", "Clotrimazol (*)", "Dùng ngoài", "Kem", "1%"),
    (S, "2. Thuốc chống nhiễm khuẩn"), ("275", "Povidon iod (*)", "Dùng ngoài", "Dung dịch", "10%"),
    (S, "3. Thuốc chống viêm, ngứa"), ("278", "Betamethason (valerat) (*)", "Dùng ngoài", "Kem, mỡ", "0,1%"),
    (S, "XVII. THUỐC ĐƯỜNG TIÊU HÓA"), (S, "1. Thuốc chống loét dạ dày, tá tràng"),
    ("309", "Magnest hydmxyd + Nhỏm hydroxyd (*)", "Uống", "Hỗn dịch", "55mg Magnesi oxyd + 64mg Nhôm oxyd/ml"),
    ("", "", "", "Viên nén, viên nhai", "200mg + 200mg, 400mg + 400mg, 500mg + 500mg"),
    ("310", "Omeprazol", "Uống", "Viên", "20mg"),
    ("311", "Pantoprazol", "Uống", "Viên", "40mg"), ("", "", "Tiêm", "Bột pha tiêm", "40mg"),
    (S, "2. Thuốc chống nôn"),
    ("313", "Metoclopramid (hydroclorid)", "Uống", "Viên", "10mg"),
    ("316", "Ondansetron (hydroclorid)", "Uống", "Viên", "4mg, 8mg, 24mg"),
    (S, "3. Thuốc chống co thắt"),
    ("317", "Alverin (citrat)", "Uống", "Viên", "10mg, 60mg"),
    ("319", "Hyoscin butylbromid", "Uống (*)", "Viên", "10mg"),
    (S, "5. Thuốc tiêu chảy"), (S, "a. Chống mất nước"),
    ("323", "Oresol (*)", "Uống", "Bột pha dung dịch", ""),
    (S, "b. Chống tiêu chảy"),
    ("324", "Atapulgit (*)", "Uống", "Bột", "3g"), ("325", "Berberin clorid (*)", "Uống", "Viên", "10mg"),
    ("326", "Loperamid (*)", "Uống", "Viên", "2mg"), ("327", "Kẽm Sulfat (*)", "Uống", "Viên", "20mg"),
    (S, "6. Thuốc điều trị bệnh trĩ"), ("328", "Diosmin (*)", "Uống", "Viên", "150mg, 300mg"),
    (S, "XVIII. HORMON, NỘI TIẾT TỐ, THUỐC TRÁNH THỤ THAI"), (S, "3. Thuốc tránh thai"),
    ("336", "Ethinylestradiol + Levonorgestrel (*)", "Uống", "Viên", "30mcg + 150mcg"),
    (S, "5. Insulin và thuốc hạ đường huyết"),
    ("345", "Glibenclamid", "Uống", "Viên", "2,5mg, 5mg"), ("346", "Gliclazid", "Uống", "Viên", "30mg, 80mg"),
    ("349", "Metformin", "Uống", "Viên", "500mg, 850mg"),
    (S, "8. Giáp trạng và kháng giáp trạng"), ("356", "Levothyroxin natri", "Uống", "Viên", "50mcg, 100mcg"),
    (S, "XXI. THUỐC DÙNG CHO MẮT, TAI, MŨI, HỌNG"), (S, "1. Thuốc chống nhiễm khuẩn, kháng virus"),
    ("403", "Cloramphenicol", "Nhỏ mắt", "Dung dịch", "0,4%"),
    ("406", "Ofloxacin", "Nhỏ mắt, tai", "Dung dịch", "0,3%"),
    ("409", "Ciprofloxacin (hydroclorid)", "Nhỏ mắt", "Dung dịch", "0,3%"), ("", "", "Tra mắt", "Mỡ", "0,3%"),
    (S, "6. Thuốc tai, mũi, họng"),
    ("418", "Naphazolin (*)", "Nhỏ mũi", "Dung dịch", "0,05%"),
    ("420", "Xylometazolin (*)", "Nhỏ mũi", "Dung dịch", "0,05%, 0,1%"),
    ("422", "Budesonid (*)", "Xịt mũi", "Dung dịch", "100mcg/Iần xịt"),
    (S, "XXV. THUỐC TÁC DỤNG TRÊN ĐƯỜNG HÔ HẤP"), (S, "1. Thuốc chữa hen và bệnh phổi tắc nghẽn mạn tính"),
    ("448", "Salbutamol (sulfat)", "Uống", "Viên", "2mg, 4mg"), ("", "", "", "Dung dịch", "0,4mg/ml"),
    ("", "", "Đường hô hấp", "Thuốc hít (khí dung)", "100mcg/lần xịt"), ("", "", "Tiêm", "Dung dịch", "50 mcg/ml"),
    ("450", "Budesonid (*)", "Đường hô hấp", "Thuốc hít (khí dung)", "100mcg/lần xịt. 200mcg/lần xịt"),
    ("452", "Ipratropium bromid", "Đường hô hấp", "Thuốc hít (khí dung)", "20 mcg/lần xịt"),
    (S, "2. Thuốc chữa rối loạn tiết dịch"),
    ("453", "Acetylcystein", "Uống (*)", "Viên", "100mg, 200mg"), ("", "", "", "Bột pha hỗn dịch", "200mg"),
    ("", "", "Tiêm", "Dung dịch", "200mg/ml"),
    ("454", "Alimemazin (*)", "Uống", "Si rô", "0,5mg/ml"),
    ("455", "Bromhexin hydroclorid (*)", "Uống", "Viên", "4mg, 8mg"),
    (S, "3. Thuốc khác"), ("456", "Dextromethorphan (*)", "Uống", "Viên", "15mg"),
    (S, "XXVII. VITAMIN VÀ CÁC CHẤT VÔ CƠ"),
    ("467", "Calci gluconat", "Uống (*)", "Viên", "500mg, 1g"), ("", "", "Tiêm", "Dung dịch", "100mg/ml"),
    ("469", "Vitamin A + D (*)", "Uống", "Viên", "5.000 IU + 500 IU"),
    ("470", "Thiamin hydroclorid hoặc nitrat (*)", "Uống", "Viên", "10mg, 50mg, 100mg"),
    ("472", "Vitamin B6 (*)", "Uống", "Viên", "25mg, 100mg"),
    ("473", "Vitamin C (*)", "Uống", "Viên", "50mg, 100mg, 500mg"),
    (S, "XXIX. THUỐC ĐIỀU TRỊ BỆNH GÚT VÀ CÁC BỆNH XƯƠNG KHỚP"), (S, "1. Thuốc điều trị bệnh gút"),
    ("478", "Allopurinol", "Uống", "Viên", "100mg, 300mg"), ("479", "Colchicin", "Uống", "Viên", "1mg"),
]


def _local_numbers():
    """STT trong nhóm (cột TT thứ hai của bản gốc): đánh lại từ 1 ở mỗi nhóm La Mã."""
    out, n = {}, 0
    for row in ROWS:
        if row[0] == S and row[1].split(".")[0].strip().isupper() and row[1][:1] in "IVXL":
            n = 0
        elif row[0] not in (S, ""):
            n += 1
            out[row[0]] = str(n)
    return out


def build_docx(path, local_first: bool = False, repeat_merged: bool = False, wide: bool = False) -> None:
    """repeat_merged: dòng tiếp của cùng thuốc LẶP LẠI STT/tên/đường dùng (như file Word chuyển từ .doc trên vbpl).
    wide: bảng 13 cột như bản thật (TT lặp 7 cột, ĐƯỜNG DÙNG lặp 3 cột)."""
    doc = Document()
    doc.add_paragraph("PHỤ LỤC I. DANH MỤC THUỐC HÓA DƯỢC, VẮC XIN, SINH PHẨM THIẾT YẾU")
    ncols = 13 if wide else 6
    t = doc.add_table(rows=0, cols=ncols)

    def widen(c):                                    # [tt, tt2, ten, duong, dang, ham] -> 13 cột như bản thật
        return [c[0]] * 6 + [c[1], c[2], c[3], c[3], c[3], c[4], c[5]] if wide else c

    def add(cells):
        cells = widen(cells)
        r = t.add_row()
        for i, v in enumerate(cells):
            r.cells[i].text = v
        return r
    add(["TT", "TT", "TÊN THUỐC", "ĐƯỜNG DÙNG", "DẠNG BÀO CHẾ", "HÀM LƯỢNG, NỒNG ĐỘ"])
    add(["(1)", "(1)", "(2)", "(3)", "(4)", "(5)"])
    for row in ROWS:
        if row[0] == S:
            r = t.add_row()
            merged = r.cells[0].merge(r.cells[ncols - 1])  # dòng tiêu đề: ô gộp ngang như bản gốc
            merged.text = row[1]
        else:
            stt, name, route, form, strength = row
            if repeat_merged:
                if stt:
                    last = {"stt": stt, "name": name, "route": route}
                else:
                    last = {"stt": last["stt"], "name": name or last["name"], "route": route or last["route"]}
                stt, name, route = last["stt"], last["name"], last["route"]
                add([stt, stt, name, route, form, strength])
                continue
            loc = LOCAL.get(stt, "") if stt else ""
            add(([loc, stt] if local_first else [stt, loc]) + [name, route, form, strength])
    doc.add_paragraph("PHỤ LỤC II. DANH MỤC THUỐC CỔ TRUYỀN THIẾT YẾU")
    t2 = doc.add_table(rows=0, cols=5)                       # phụ lục II: đánh số lại từ đầu
    for vals in (["TT", "TT", "TÊN THUỐC", "ĐƯỜNG DÙNG", ""], ["14", "6", "Actiso, Thổ phục linh", "Uống", ""]):
        r = t2.add_row()
        for i, v in enumerate(vals):
            r.cells[i].text = v
    doc.save(str(path))


def build_html(path, rowspan: bool = True) -> None:
    """Trang toàn văn dạng HTML (như lưu bằng Ctrl+S trên trình duyệt). rowspan=True: ô TT/tên/đường dùng
    gộp dọc bằng rowspan (ô bị gộp không có trong HTML); False: ô trống như bản .docx."""
    import html as h
    data = [r for r in ROWS]
    out = ['<html><head><meta charset="utf-8"><title>Thông tư 19/2018/TT-BYT</title></head><body>',
           '<table><tr><td>Bảng bố cục trang (không phải phụ lục)</td></tr></table>',
           "<p>PHỤ LỤC I</p><table border=1>",
           "<tr><th>TT</th><th>TT</th><th>TÊN<br>THUỐC</th><th>ĐƯỜNG DÙNG</th><th>DẠNG BÀO CHẾ</th>"
           "<th>HÀM LƯỢNG, NỒNG ĐỘ</th></tr>",
           "<tr>" + "".join(f"<td>({n})</td>" for n in (1, 1, 2, 3, 4, 5)) + "</tr>"]
    covered = set()                                     # (dòng, cột) đã bị rowspan che
    for i, row in enumerate(data):
        if row[0] == S:
            out.append(f'<tr><td colspan="6"><b>{h.escape(row[1])}</b></td></tr>')
            continue
        stt, name, route, form, strength = row
        vals = [stt, LOCAL.get(stt, "") if stt else "", name, route, form, strength]
        cells = []
        for c, v in enumerate(vals):
            if (i, c) in covered:
                continue
            span = 1
            if rowspan and c <= 3 and v:
                j = i + 1
                while j < len(data) and data[j][0] != S and data[j][0] == "" and \
                        (c <= 1 or (data[j][c - 1] == "")):
                    covered.add((j, c))
                    span += 1
                    j += 1
            attr = f' rowspan="{span}"' if span > 1 else ""
            cells.append(f"<td{attr}>{h.escape(v)}</td>")
        out.append("<tr>" + "".join(cells) + "</tr>")
    out.append("</table><p>PHỤ LỤC II</p><table><tr><td>TT</td><td>TÊN THUỐC</td><td>ĐƯỜNG DÙNG</td>"
               "<td>DẠNG BÀO CHẾ</td><td>HÀM LƯỢNG</td></tr><tr><td>1</td><td>Actiso</td><td>Uống</td>"
               "<td>Viên</td><td>x</td></tr></table></body></html>")
    with open(path, "w", encoding="utf-8") as f:
        f.write("\n".join(out))


LOCAL = _local_numbers()
